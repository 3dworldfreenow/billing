"""GST return workbook for the accountant: GSTR-1 tables (B2B, B2CS, credit notes, HSN
summary, documents issued) and a GSTR-3B summary, for intra-state sales (CGST + SGST)."""

import io
from collections import defaultdict
from decimal import Decimal

from openpyxl import Workbook
from openpyxl.styles import Font

from .models import D, Purchase, PurchaseItem, Sale, Setting, db, money

UQC = {"kg": "KGS", "g": "GMS", "ltr": "LTR", "ml": "MLT", "box": "BOX", "dozen": "DOZ", "pack": "PAC"}


def _sheet(wb, title, header, rows):
    ws = wb.create_sheet(title)
    ws.append(header)
    for cell in ws[1]:
        cell.font = Font(bold=True)
    for row in rows:
        ws.append([float(v) if isinstance(v, Decimal) else v for v in row])
    for column in ws.columns:
        ws.column_dimensions[column[0].column_letter].width = max(12, min(40, len(str(column[0].value)) + 4))


def gst_workbook(start, end):
    store = Setting.get()
    state = store.state_code or ""
    docs = Sale.query.filter(Sale.status.in_(("saved", "cancelled")), Sale.created_at >= start,
                             Sale.created_at < end).order_by(Sale.id).all()
    saved = [s for s in docs if s.status == "saved" and s.taxed]

    b2b, cdnr = [], []
    b2cs = defaultdict(lambda: [Decimal(0)] * 3)
    hsn = defaultdict(lambda: [None, None, Decimal(0), Decimal(0), Decimal(0), Decimal(0), Decimal(0)])
    out_taxable = out_cgst = out_sgst = Decimal(0)
    for sale in saved:
        registered = sale.party is not None and bool(sale.party.gstin)
        place = sale.party.gstin[:2] if registered else state
        for rate, (taxable, cgst, sgst) in sale.tax_summary().items():
            out_taxable, out_cgst, out_sgst = out_taxable + taxable, out_cgst + cgst, out_sgst + sgst
            if registered and not sale.is_return:
                b2b.append([sale.party.gstin, sale.party.name, sale.number, sale.created_at.strftime("%d-%m-%Y"),
                            D(sale.net), place, "N", "Regular", rate, taxable, cgst, sgst, 0])
            elif registered:
                cdnr.append([sale.party.gstin, sale.party.name, sale.doc_no, sale.created_at.strftime("%d-%m-%Y"),
                             "C", place, -D(sale.net), rate, -taxable, -cgst, -sgst,
                             sale.return_of.number if sale.return_of else ""])
            else:  # unregistered: one line per rate, credit notes netted off
                row = b2cs[rate]
                row[0], row[1], row[2] = row[0] + taxable, row[1] + cgst, row[2] + sgst
        for item in sale.items:
            code = (item.product.hsn if item.product else None) or ""
            unit = item.product.unit if item.product else "pcs"
            row = hsn[(code, D(item.gst_rate))]
            row[0] = row[0] or item.name
            row[1] = UQC.get(unit, "NOS")
            row[2] += D(item.qty)
            row[3] += D(item.amount)
            tax = money(D(item.revenue) * D(item.gst_rate) / 100)
            row[4] += D(item.revenue)
            row[5] += money(tax / 2)
            row[6] += tax - money(tax / 2)

    def series(numbers):
        numbers = [n for n in numbers if n]
        return (min(numbers), max(numbers), len(numbers)) if numbers else ("", "", 0)

    bills = [s for s in docs if not s.is_return]
    notes = [s for s in docs if s.is_return]
    doc_rows = []
    for label, group, attr, prefix in (("Invoices for outward supply", bills, "number", ""),
                                       ("Credit notes", notes, "return_number", "CN-")):
        first, last, total = series([getattr(s, attr) for s in group])
        cancelled = sum(1 for s in group if s.status == "cancelled")
        if total:
            doc_rows.append([label, f"{prefix}{first}", f"{prefix}{last}", total, cancelled, total - cancelled])

    purchases = db.session.query(PurchaseItem.taxable, PurchaseItem.tax).join(Purchase).filter(
        Purchase.status == "saved", Purchase.created_at >= start, Purchase.created_at < end).all()
    in_taxable = sum((D(t) for t, _ in purchases), Decimal(0))
    in_tax = sum((D(x) for _, x in purchases), Decimal(0))
    in_cgst = money(in_tax / 2)

    wb = Workbook()
    wb.remove(wb.active)
    _sheet(wb, "GSTR-3B summary", ["Section", "Taxable value", "CGST", "SGST", "Total tax"], [
        ["3.1(a) Outward taxable supplies", out_taxable, out_cgst, out_sgst, out_cgst + out_sgst],
        ["4(A)(5) Input tax credit (purchases)", in_taxable, in_cgst, in_tax - in_cgst, in_tax],
        ["Net tax payable (before other adjustments)", "", out_cgst - in_cgst, out_sgst - (in_tax - in_cgst),
         out_cgst + out_sgst - in_tax],
    ])
    _sheet(wb, "b2b", ["GSTIN of recipient", "Receiver name", "Invoice number", "Invoice date", "Invoice value",
                       "Place of supply", "Reverse charge", "Invoice type", "Rate", "Taxable value", "CGST", "SGST",
                       "Cess"], b2b)
    _sheet(wb, "b2cs", ["Type", "Place of supply", "Rate", "Taxable value", "CGST", "SGST", "Cess"],
           [["OE", state, rate, *values, 0] for rate, values in sorted(b2cs.items())])
    _sheet(wb, "cdnr", ["GSTIN of recipient", "Receiver name", "Note number", "Note date", "Note type",
                        "Place of supply", "Note value", "Rate", "Taxable value", "CGST", "SGST",
                        "Original invoice"], cdnr)
    _sheet(wb, "hsn", ["HSN", "Description", "UQC", "Total quantity", "Total value", "Rate", "Taxable value",
                       "CGST", "SGST"],
           [[code, name, uqc, qty, value, rate, taxable, cgst, sgst]
            for (code, rate), (name, uqc, qty, value, taxable, cgst, sgst) in sorted(hsn.items())])
    _sheet(wb, "docs", ["Nature of document", "Sr. no. from", "Sr. no. to", "Total number", "Cancelled",
                        "Net issued"], doc_rows)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
