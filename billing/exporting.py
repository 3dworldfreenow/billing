"""Excel/CSV export of report tables, and reading uploaded spreadsheets."""

import csv
import io
from datetime import date, datetime
from decimal import Decimal

from flask import Response
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font

XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def _plain(cell):
    return cell["v"] if isinstance(cell, dict) else cell


def export_response(name, columns, rows, totals=None, fmt="xlsx"):
    """columns: [(label, kind)], rows: lists of values (or {"v": value, "url": ...})."""
    data = [[_plain(c) for c in row] for row in rows]
    if totals:
        data.append([_plain(c) for c in totals])
    labels = [label for label, _ in columns]
    stamp = date.today().isoformat()

    if fmt == "csv":
        out = io.StringIO()
        writer = csv.writer(out)
        writer.writerow(labels)
        writer.writerows(["" if v is None else v for v in row] for row in data)
        body, mimetype, ext = "﻿" + out.getvalue(), "text/csv", "csv"  # BOM so Excel reads ₹/Hindi text
    else:
        wb = Workbook()
        ws = wb.active
        ws.title = name[:31]
        ws.append(labels)
        for cell in ws[1]:
            cell.font = Font(bold=True)
        for row in data:
            ws.append([float(v) if isinstance(v, Decimal) else v for v in row])
        if totals:
            for cell in ws[ws.max_row]:
                cell.font = Font(bold=True)
        for i, (label, kind) in enumerate(columns, start=1):
            letter = ws.cell(row=1, column=i).column_letter
            ws.column_dimensions[letter].width = max(12, len(label) + 2) if kind != "text" else 32
            if kind == "money":
                for cell in ws[letter][1:]:
                    cell.number_format = "#,##0.00"
        buf = io.BytesIO()
        wb.save(buf)
        body, mimetype, ext = buf.getvalue(), XLSX, "xlsx"

    filename = f"{name.lower().replace(' ', '-')}-{stamp}.{ext}"
    return Response(body, mimetype=mimetype, headers={"Content-Disposition": f'attachment; filename="{filename}"'})


def read_rows(upload):
    """Rows of an uploaded .xlsx or .csv file as dicts keyed by the header row."""
    filename = (upload.filename or "").lower()
    if filename.endswith(".xlsx"):
        ws = load_workbook(upload, read_only=True, data_only=True).active
        rows = ws.iter_rows(values_only=True)
        header = [str(h or "").strip() for h in next(rows, [])]
        for row in rows:
            if any(v not in (None, "") for v in row):
                yield {h: v for h, v in zip(header, row)}
    elif filename.endswith(".csv"):
        text = upload.read().decode("utf-8-sig", errors="replace")
        yield from csv.DictReader(io.StringIO(text))
    else:
        raise ValueError("Upload an .xlsx or .csv file.")


def cell_text(value):
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        value = int(value)  # Excel turns barcodes into floats
    if isinstance(value, (date, datetime)):
        return value.isoformat()
    return str(value).strip()
