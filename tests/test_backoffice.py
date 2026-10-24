import io
from decimal import Decimal

from openpyxl import Workbook, load_workbook

from billing.models import Party, Product, Purchase, TaxSlab, User, db
from conftest import new_sale


def test_purchase_adds_stock_and_updates_prices(client):
    supplier = Party(kind="supplier", name="PRK Commodities")
    db.session.add(supplier)
    db.session.commit()
    resp = client.get("/purchases/new")
    purchase_id = int(resp.headers["Location"].rstrip("/").rsplit("/", 1)[1])
    client.post(f"/purchases/{purchase_id}/scan", data={"code": "8904018301680"})
    item = db.session.get(Purchase, purchase_id).items[0]
    client.post(f"/purchases/{purchase_id}/update", data={
        "party_id": supplier.id, "payment_mode": "credit", f"qty_{item.id}": "10", f"free_qty_{item.id}": "2",
        f"rate_{item.id}": "36", f"disc_pct_{item.id}": "0", f"gst_rate_{item.id}": "18",
        f"mrp_{item.id}": "48", f"sale_rate_{item.id}": "45", "action": "save",
    })
    purchase = db.session.get(Purchase, purchase_id)
    assert purchase.status == "saved"
    assert purchase.net == Decimal("425")  # 360 + 64.80 GST, rounded
    soap = Product.query.filter_by(barcode="8904018301680").one()
    assert soap.stock == Decimal("22")
    assert soap.purchase_rate == Decimal("30.00")  # 360 spread over 12 units
    assert (soap.mrp, soap.sale_rate) == (Decimal("48.00"), Decimal("45.00"))
    assert supplier.balance == Decimal("425.00")


def test_unknown_purchase_barcode_creates_product_and_adds_it(client):
    purchase_id = int(client.get("/purchases/new").headers["Location"].rsplit("/", 1)[1])
    resp = client.post(f"/purchases/{purchase_id}/scan", data={"code": "8909999999999"}, follow_redirects=True)
    assert b"Create product with barcode 8909999999999" in resp.data
    client.post(f"/products/new?barcode=8909999999999&add_to=purchase:{purchase_id}",
                data={"name": "New Biscuit", "barcode": "8909999999999", "mrp": "10", "sale_rate": "10",
                      "gst_rate": "5"})
    assert db.session.get(Purchase, purchase_id).items[0].product.name == "New Biscuit"


def test_product_import_swilerp_style_and_export(client):
    wb = Workbook()
    ws = wb.active
    ws.append(["NameToDisplay", "Curr.Qty", "GroupName", "Barcode", "Category", "MRP", "SaleRate", "GST %"])
    ws.append(["24 CARATS CHICKEN MASALA 100 GM", 15, "PRK COMMODITIES", 904341101179, "MASALA", 60, 58, 5])
    ws.append(["Medimix Soap 125g", 99, "", 8904018301680, "SOAP", 46, 44, 18])  # existing -> update
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    resp = client.post("/products/import", data={"file": (buf, "products.xlsx")},
                       content_type="multipart/form-data", follow_redirects=True)
    assert b"1 new, 1 updated" in resp.data
    masala = Product.query.filter_by(barcode="904341101179").one()
    assert (masala.stock, masala.brand, masala.category) == (Decimal("15"), "PRK COMMODITIES", "MASALA")
    soap = Product.query.filter_by(barcode="8904018301680").one()
    assert soap.mrp == Decimal("46") and soap.stock == Decimal("10")  # stock kept unless asked

    xlsx = client.get("/products/export?format=xlsx")
    rows = list(load_workbook(io.BytesIO(xlsx.data)).active.values)
    assert rows[0][:2] == ("Name", "Barcode") and len(rows) == 5
    assert b"Name,Barcode" in client.get("/products/export?format=csv").data


def test_reports_render_and_export(client):
    sale_id = new_sale(client)
    client.post(f"/sales/{sale_id}/scan", data={"code": "8904018301680"})
    client.post(f"/sales/{sale_id}/update", data={"action": "save"})
    for url in ["/reports/", "/reports/sales", "/reports/daily", "/reports/products", "/reports/purchases",
                "/reports/gst", "/reports/stock", "/reports/stock?low=1", "/reports/dues"]:
        assert client.get(url).status_code == 200, url
    gst = client.get("/reports/gst?export=csv").data.decode()
    assert "18%,35.59,3.21,3.20" in gst
    assert client.get("/reports/sales?export=xlsx").mimetype.endswith("sheet")


def test_tax_slabs_and_settings(client):
    assert {float(s.rate) for s in TaxSlab.query} == {0, 5, 18, 40, 3, 0.25}
    client.post("/settings/slabs", data={"rate": "12", "name": "GST 12% (old)"})
    slab = TaxSlab.query.filter_by(rate=12).one()
    client.post(f"/settings/slabs/{slab.id}/toggle")
    assert not slab.active
    assert b"already exists" in client.post("/settings/slabs", data={"rate": "5"}, follow_redirects=True).data


def test_stock_adjustment(client):
    soap = Product.query.filter_by(barcode="8904018301680").one()
    client.post(f"/stock/{soap.id}/adjust", data={"reason": "count", "actual": "7"})
    client.post(f"/stock/{soap.id}/adjust", data={"reason": "expired", "lost": "2"})
    assert soap.stock == Decimal("5")
    assert b"Expired" in client.get("/stock/").data


def test_login_and_permissions(app):
    client = app.test_client()
    assert client.get("/pos").headers["Location"].endswith("/setup")
    client.post("/setup", data={"name": "Owner", "username": "admin", "password": "secret123"})
    client.post("/users/new", data={"name": "Cashier", "username": "ravi", "password": "cash1234",
                                    "perm_sales": "1", "active": "1"})
    assert User.query.filter_by(username="ravi").one().permission_set == {"sales"}
    client.post("/logout")
    assert "/login" in client.get("/pos").headers["Location"]
    assert b"Wrong username" in client.post("/login", data={"username": "ravi", "password": "x"},
                                            follow_redirects=True).data
    resp = client.post("/login", data={"username": "ravi", "password": "cash1234"})
    assert resp.headers["Location"].endswith("/pos")
    assert client.get("/pos").status_code == 302
    for url in ["/reports/", "/products/", "/settings/", "/users"]:
        assert client.get(url).status_code == 403, url


def test_pages_render(client):
    for url in ["/products/", "/products/new", "/products/import", "/customers", "/suppliers", "/customers/new",
                "/suppliers/new", "/purchases/", "/sales", "/stock/", "/stock/?q=soap", "/settings/", "/users",
                "/users/new", "/account"]:
        assert client.get(url).status_code == 200, url


def test_each_cashier_gets_their_own_bill(app, client):
    client.post("/users/new", data={"name": "Counter 2", "username": "counter2", "password": "cash1234",
                                    "perm_sales": "1", "active": "1"})
    other = app.test_client()
    other.post("/login", data={"username": "counter2", "password": "cash1234"})
    assert new_sale(client) != new_sale(other)


def test_live_search_api_and_lists(client):
    names = [p["name"] for p in client.get("/api/products?q=sal").get_json()]
    assert names[0] == "Tata Salt 1kg"  # name starting with the text ranks first
    assert [p["name"] for p in client.get("/api/products?q=soap 125").get_json()] == ["Medimix Soap 125g"]
    assert client.get("/api/products?q=8904018").get_json()[0]["barcode"] == "8904018301680"  # barcode start
    client.post("/customers/new", data={"name": "Ramesh Kumar", "phone": "9876543210"})
    assert client.get("/api/customers?q=98765").get_json()[0]["name"] == "Ramesh Kumar"
    assert b"Tata Salt" in client.get("/products/?q=salt").data
    assert b"Medimix" not in client.get("/products/?q=salt").data
    soap = Product.query.filter_by(barcode="8904018301680").one()
    assert b"System stock" in client.get(f"/stock/?product_id={soap.id}").data


def test_search_api_needs_permission(app, client):
    client.post("/users/new", data={"name": "Clerk", "username": "clerk", "password": "clerk123",
                                    "perm_parties": "1", "active": "1"})
    other = app.test_client()
    other.post("/login", data={"username": "clerk", "password": "clerk123"})
    assert other.get("/api/products?q=salt").status_code == 403
    assert other.get("/api/customers?q=a").status_code == 200
