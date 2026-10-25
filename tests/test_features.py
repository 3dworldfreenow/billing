import io
import json
import socket
import threading
from datetime import date, timedelta
from decimal import Decimal

from openpyxl import Workbook, load_workbook

from billing.models import (AuditLog, Batch, Party, Product, Promotion, Purchase, Sale, Setting, User, db)
from billing.services import ean13_check
from conftest import new_sale

SOAP, SALT = "8904018301680", "8901058000290"


def product(barcode):
    return Product.query.filter_by(barcode=barcode).one()


def sell(client, *codes, **form):
    sale_id = new_sale(client)
    for code in codes:
        client.post(f"/sales/{sale_id}/scan", data={"code": code})
    client.post(f"/sales/{sale_id}/update", data={"action": "save", **form})
    return db.session.get(Sale, sale_id)


def purchase(client, supplier, lines, **form):
    purchase_id = int(client.get("/purchases/new").headers["Location"].rsplit("/", 1)[1])
    for code in lines:
        client.post(f"/purchases/{purchase_id}/scan", data={"code": code})
    doc = db.session.get(Purchase, purchase_id)
    data = {"party_id": supplier.id, "payment_mode": "credit", "action": "save"}
    for item in doc.items:
        for key, value in lines[item.product.barcode].items():
            data[f"{key}_{item.id}"] = value
    data.update(form)
    client.post(f"/purchases/{purchase_id}/update", data=data)
    return db.session.get(Purchase, purchase_id)


def supplier():
    party = Party(kind="supplier", name="PRK Commodities", phone="9000000001")
    db.session.add(party)
    db.session.commit()
    return party


# --- Batches, expiry, multiple MRP ---------------------------------------------------


def test_batches_fefo_mrp_switch_and_cancel(client):
    soap = product(SOAP)  # 10 in stock at MRP 45
    expiry = date.today() + timedelta(days=200)
    purchase(client, supplier(), {SOAP: {"qty": "12", "rate": "32", "mrp": "48", "sale_rate": "46", "gst_rate": "18",
                                         "batch_no": "B77", "expiry": expiry.isoformat()}})
    assert [(b.batch_no, b.qty, b.mrp) for b in soap.batches] == [("OPENING", 10, Decimal("45.00")),
                                                                    ("B77", 12, Decimal("48.00"))]
    assert soap.stock == 22 and soap.batches[1].expiry == expiry
    # The batch without expiry sells after the one that expires: B77 first.
    sale_id = new_sale(client)
    client.post(f"/sales/{sale_id}/scan", data={"code": SOAP})
    item = db.session.get(Sale, sale_id).items[0]
    assert (item.batch.batch_no, item.mrp, item.rate) == ("B77", Decimal("48.00"), Decimal("46.00"))
    # Cashier switches the line to the old MRP 45 stock.
    opening = soap.batches[0]
    client.post(f"/sales/{sale_id}/update", data={f"batch_{item.id}": opening.id, "action": "save"})
    sale = db.session.get(Sale, sale_id)
    assert (sale.items[0].mrp, sale.items[0].rate, sale.net) == (Decimal("45.00"), Decimal("42.00"), Decimal("42"))
    assert (opening.qty, soap.stock) == (9, 21)
    client.post(f"/sales/{sale_id}/cancel")
    assert (opening.qty, soap.stock) == (10, 22)
    assert b"Batches" in client.get(f"/products/{soap.id}").data


def test_new_mrp_without_batch_number_still_keeps_old_stock_apart(client):
    salt = product(SALT)  # 50 at MRP 28
    purchase(client, supplier(), {SALT: {"qty": "20", "rate": "24", "mrp": "30", "sale_rate": "30", "gst_rate": "5"}})
    assert sorted((b.mrp, b.qty) for b in salt.batches) == [(Decimal("28.00"), 50), (Decimal("30.00"), 20)]
    assert salt.stock == sum(b.qty for b in salt.batches)


def test_stock_adjust_needs_batch_for_batched_product(client):
    soap = product(SOAP)
    purchase(client, supplier(), {SOAP: {"qty": "5", "rate": "32", "mrp": "48", "gst_rate": "18", "batch_no": "B1"}})
    resp = client.post(f"/stock/{soap.id}/adjust", data={"reason": "count", "actual": "1"}, follow_redirects=True)
    assert b"Choose the batch" in resp.data
    b1 = next(b for b in soap.batches if b.batch_no == "B1")
    client.post(f"/stock/{soap.id}/adjust", data={"reason": "expired", "lost": "2", "batch_id": b1.id})
    assert (b1.qty, soap.stock) == (3, 13)


def test_expiry_report_and_dashboard(client):
    soap = product(SOAP)
    db.session.add(Batch(product=soap, batch_no="OLD", expiry=date.today() - timedelta(days=1), mrp=45, qty=2))
    db.session.commit()
    assert b"EXPIRED" in client.get("/reports/expiry").data
    assert b"Expiring in 30 days" in client.get("/reports/").data


# --- Returns / credit notes ------------------------------------------------------------


def test_return_makes_credit_note_and_restores_stock(client):
    sale = sell(client, f"4*{SOAP}", discount="8")  # 168 - 8 = 160
    item = sale.items[0]
    resp = client.post(f"/sales/{sale.id}/return", data={f"qty_{item.id}": "1", "mode": "cash"})
    note = Sale.query.filter_by(kind="return").one()
    assert resp.headers["Location"].endswith(f"/sales/{note.id}")
    assert (note.doc_no, note.net, note.return_of_id) == ("CN-1", Decimal("-40"), sale.id)  # 42 less 1/4 of 8
    assert product(SOAP).stock == 7
    assert b"at most 3" in client.post(f"/sales/{sale.id}/return", data={f"qty_{item.id}": "5"},
                                       follow_redirects=True).data
    assert b"Cancel the credit note first" in client.post(f"/sales/{sale.id}/cancel", follow_redirects=True).data
    from billing.views.sales import bill_text
    text = bill_text(note)
    assert "Credit Note: CN-1" in text and "REFUND: ₹40.00" in text and "Refunded by CASH" in text
    assert b"CREDIT NOTE" in client.get(f"/sales/{note.id}/receipt").data
    assert AuditLog.query.filter_by(action="return").count() == 1


def test_refund_to_account_reduces_customer_balance(client):
    client.post("/customers/new", data={"name": "Ramesh", "phone": "9876543210"})
    customer = Party.query.filter_by(name="Ramesh").one()
    sale = sell(client, f"2*{SALT}", party_id=customer.id, payment_mode="credit")
    assert customer.balance == Decimal("56.00")
    client.post(f"/sales/{sale.id}/return", data={f"qty_{sale.items[0].id}": "1", "mode": "credit"})
    assert customer.balance == Decimal("28.00")


# --- Scale barcodes and labels -----------------------------------------------------------


def test_weighing_scale_barcode_adds_weight(client):
    client.post("/settings/devices", data={"scale_prefix": "21", "scale_item_digits": "5", "scale_value": "weight",
                                           "printer_mode": "browser", "printer_width": "48"})
    sugar = Product.query.filter_by(name="Loose Sugar").one()
    sugar.plu = "101"
    db.session.commit()
    body = "21" + "00101" + "01250"
    sale_id = new_sale(client)
    client.post(f"/sales/{sale_id}/scan", data={"code": body + ean13_check(body)})
    item = db.session.get(Sale, sale_id).items[0]
    assert (item.name, item.qty, item.amount) == ("Loose Sugar", Decimal("1.250"), Decimal("56.25"))


def test_price_scale_barcode(client):
    store = Setting.get()
    store.scale_prefix, store.scale_item_digits, store.scale_value = "2", 5, "price"
    Product.query.filter_by(name="Loose Sugar").one().plu = "7"
    db.session.commit()
    body = "2" + "00007" + "009000"  # ₹90.00 at ₹45/kg -> 2 kg
    sale_id = new_sale(client)
    client.post(f"/sales/{sale_id}/scan", data={"code": body + ean13_check(body)})
    assert db.session.get(Sale, sale_id).items[0].qty == Decimal("2.000")


def test_generate_barcode_and_print_labels(client):
    sugar = Product.query.filter_by(name="Loose Sugar").one()
    client.post(f"/products/{sugar.id}/barcode")
    assert len(sugar.barcode) == 13 and sugar.barcode.startswith("200") and ean13_check(sugar.barcode[:12]) == \
        sugar.barcode[12]
    client.post("/products/labels", data={"product_id": sugar.id, "copies": "3"})
    client.post("/products/labels", data={"code": SOAP, "copies": "2"})
    assert b"Loose Sugar" in client.get("/products/labels").data
    page = client.get("/products/labels/print?layout=a4-65").data.decode()
    assert page.count('class="label"') == 5 and page.count("<svg version") == 5  # barcodes, not icons


def test_labels_from_purchase(client):
    doc = purchase(client, supplier(), {SALT: {"qty": "6", "rate": "22", "gst_rate": "5"}})
    client.post(f"/products/labels/from-purchase/{doc.id}")
    page = client.get("/products/labels/print?layout=roll-50x25").data.decode()
    assert page.count('class="label"') == 6


# --- Offers, loyalty, wholesale, limits ----------------------------------------------------


def test_buy_two_get_one_and_bill_offer(client):
    soap = product(SOAP)
    db.session.add_all([
        Promotion(name="Soap 2+1", kind="bxgy", product=soap, buy_qty=2, free_qty=1),
        Promotion(name="5% off above 100", kind="bill", percent=5, min_amount=100),
        Promotion(name="Expired", kind="percent", percent=50, ends_on=date.today() - timedelta(days=1)),
    ])
    db.session.commit()
    sale = sell(client, f"4*{SOAP}")  # 168, one soap free (42) -> 126 (>= 100); bill offer 5% of 126 = 6.30
    item = sale.items[0]
    assert (item.promo_note, item.promo_amount, item.amount) == ("Soap 2+1", Decimal("42.00"), Decimal("126.00"))
    assert (sale.promo_note, sale.promo_discount, sale.net) == ("5% off above 100", Decimal("6.30"), Decimal("120"))


def test_category_percent_offer_with_happy_hours(client):
    from datetime import datetime, time
    product(SALT).category = "SALT"
    db.session.add(Promotion(name="Salt 10%", kind="percent", category="salt", percent=10,
                             start_time=time(0, 0), end_time=time(23, 59, 59)))
    db.session.commit()
    sale = sell(client, SALT)
    assert sale.items[0].amount == Decimal("25.20")
    off_hours = Promotion.running(datetime.combine(date.today(), time(0, 0)) - timedelta(seconds=1))
    assert [p.name for p in off_hours] == ["Salt 10%"] or all(p.name != "Salt 10%" for p in off_hours)


def test_loyalty_earn_and_redeem(client):
    client.post("/offers/loyalty", data={"loyalty_rate": "1", "loyalty_value": "1", "loyalty_min_redeem": "0"})
    client.post("/customers/new", data={"name": "Sita", "phone": "9123456780"})
    sita = Party.query.filter_by(name="Sita").one()
    first = sell(client, f"25*{SOAP}", party_id=sita.id)  # ₹1050 -> 10 points
    assert first.points_earned == 10 and sita.points == 10
    resp = client.post(f"/sales/{new_sale(client)}/update", data={"party_id": sita.id, "points": "50"},
                       follow_redirects=True)
    assert b"has 10 points" in resp.data
    second = sell(client, SALT, party_id=sita.id, points="10")
    assert (second.redeem_amount, second.net, second.points_redeemed) == (Decimal("10.00"), Decimal("18"), 10)
    assert sita.points == 0


def test_wholesale_customer_gets_wholesale_rate(client):
    salt = product(SALT)
    salt.wholesale_rate = Decimal("25")
    db.session.commit()
    client.post("/customers/new", data={"name": "Kirana Store", "wholesale": "1"})
    shop = Party.query.filter_by(name="Kirana Store").one()
    assert shop.price_list == "wholesale"
    sale_id = new_sale(client)
    client.post(f"/sales/{sale_id}/scan", data={"code": SALT})
    assert db.session.get(Sale, sale_id).items[0].rate == Decimal("28.00")
    client.post(f"/sales/{sale_id}/update", data={"party_id": shop.id, "action": "update"})
    assert db.session.get(Sale, sale_id).items[0].rate == Decimal("25.00")


def test_cashier_discount_limit_rates_and_cancel(app, client):
    client.post("/users/new", data={"name": "Ravi", "username": "ravi", "password": "cash1234", "perm_sales": "1",
                                    "max_discount": "5", "active": "1"})
    cashier = app.test_client()
    cashier.post("/login", data={"username": "ravi", "password": "cash1234"})
    sale_id = new_sale(cashier)
    cashier.post(f"/sales/{sale_id}/scan", data={"code": SOAP})
    item = db.session.get(Sale, sale_id).items[0]
    assert b"discount limit is 5%" in cashier.post(f"/sales/{sale_id}/update", data={f"disc_{item.id}": "10"},
                                                   follow_redirects=True).data
    assert b"not allowed to change rates" in cashier.post(f"/sales/{sale_id}/update", data={f"rate_{item.id}": "1"},
                                                          follow_redirects=True).data
    assert b"discount limit is 5% of the bill" in cashier.post(
        f"/sales/{sale_id}/update", data={"discount": "10"}, follow_redirects=True).data
    cashier.post(f"/sales/{sale_id}/update", data={f"disc_{item.id}": "5", "action": "save"})
    assert db.session.get(Sale, sale_id).status == "saved"
    assert cashier.post(f"/sales/{sale_id}/cancel").status_code == 403
    assert cashier.get(f"/sales/{sale_id}/return").status_code == 403
    assert cashier.get("/cash/").status_code == 403


def test_audit_log_records_cancel_and_price_change(client):
    sale = sell(client, SOAP)
    client.post(f"/sales/{sale.id}/cancel")
    soap = product(SOAP)
    client.post(f"/products/{soap.id}", data={"name": soap.name, "barcode": SOAP, "mrp": "45", "sale_rate": "40",
                                              "purchase_rate": "30", "gst_rate": "18", "stock": "10"})
    page = client.get("/settings/audit").data.decode()
    assert "cancel bill" in page and "sale rate 42 → 40" in page


# --- Cash register, expenses, profit ---------------------------------------------------------


def test_day_close_expected_cash(client):
    client.post("/cash/entry", data={"kind": "opening", "amount": "1000"})
    sell(client, f"4*{SOAP}")  # 168 cash
    sell(client, SALT, payment_mode="upi")
    client.post("/cash/expense", data={"category": "Tea & snacks", "amount": "50", "mode": "cash"})
    client.post("/cash/entry", data={"kind": "out", "amount": "100", "note": "owner"})
    page = client.get("/cash/").data.decode()
    assert "₹1,018.00" in page and "₹28.00" in page  # expected cash, UPI
    resp = client.post("/cash/close", data={"day": date.today().isoformat(), "counted": "1000"}, follow_redirects=True)
    assert "Short ₹18.00" in resp.data.decode()
    profit = client.get("/reports/profit?export=csv").data.decode()
    assert "Expense: Tea & snacks,-50" in profit and "NET PROFIT" in profit


# --- Purchase orders and reorder -----------------------------------------------------------------


def test_reorder_order_receive_and_whatsapp(client):
    party = supplier()
    soap = product(SOAP)
    soap.min_stock = Decimal("12")  # stock 10 is below the alert level
    db.session.commit()
    purchase(client, party, {SALT: {"qty": "1", "rate": "22", "gst_rate": "5"}})
    page = client.get("/purchases/reorder").data.decode()
    assert "Medimix Soap" in page and "Tata Salt" not in page
    resp = client.post("/purchases/reorder", data={"party_id": "", f"pick_{soap.id}": "1", f"qty_{soap.id}": "14"})
    order = Purchase.query.filter_by(status="order").one()
    assert order.order_number == 1 and order.items[0].qty == 14 and resp.status_code == 302
    assert soap.stock == 10  # an order doesn't change stock
    wa = client.get(f"/purchases/{order.id}/whatsapp").headers["Location"]
    assert "Purchase%20Order%201" in wa
    client.post(f"/purchases/{order.id}/update", data={"party_id": party.id, "action": "receive"})
    assert db.session.get(Purchase, order.id).status == "draft"
    client.post(f"/purchases/{order.id}/update", data={"party_id": party.id, "action": "save"})
    assert db.session.get(Purchase, order.id).status == "saved" and soap.stock == 24


# --- GST returns and reports ---------------------------------------------------------------------


def test_gstr1_workbook(client):
    Setting.get().gstin = "08ABCDE1234F1Z5"
    client.post("/customers/new", data={"name": "Hotel Raj", "gstin": "08AAACH1234A1Z9"})
    hotel = Party.query.filter_by(name="Hotel Raj").one()
    b2b = sell(client, f"2*{SOAP}", party_id=hotel.id)
    sell(client, SOAP, SALT)
    walkin = sell(client, SOAP)
    client.post(f"/sales/{walkin.id}/cancel")
    wb = load_workbook(io.BytesIO(client.get("/reports/gstr1").data))
    assert wb.sheetnames == ["GSTR-3B summary", "b2b", "b2cs", "cdnr", "hsn", "docs"]
    b2b_rows = list(wb["b2b"].values)[1:]
    assert b2b_rows == [("08AAACH1234A1Z9", "Hotel Raj", b2b.number, b2b.created_at.strftime("%d-%m-%Y"), 84, "08", "N",
                         "Regular", 18, 71.19, 6.41, 6.40, 0)]
    b2cs = {row[2]: row for row in list(wb["b2cs"].values)[1:]}
    assert b2cs[18][3] == 35.59 and b2cs[0][3] == 28
    docs = list(wb["docs"].values)[1]
    assert docs[3:] == (3, 1, 2)  # 3 invoices, 1 cancelled


def test_new_reports_render_and_export(client):
    sell(client, SOAP, SALT)
    for url in ["/reports/profit", "/reports/expenses", "/reports/categories", "/reports/movement",
                "/reports/movement?show=dead", "/reports/expiry?days=0", "/reports/gst"]:
        assert client.get(url).status_code == 200, url
    movement = client.get("/reports/movement?export=csv").data.decode()
    assert "Medimix Soap 125g" in movement and "Loose Sugar,,Dead" in movement


# --- Parties ------------------------------------------------------------------------------------------


def test_import_customers_with_opening_balances_and_reminder(client):
    wb = Workbook()
    ws = wb.active
    ws.append(["Customer Name", "Mobile No", "Opening Balance"])
    ws.append(["Mohan", 9876500001, "1,200 Dr"])
    ws.append(["Sohan", 9876500002, "300 Cr"])
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    resp = client.post("/customers/import", data={"file": (buf, "customers.xlsx")}, content_type="multipart/form-data",
                       follow_redirects=True)
    assert b"2 new, 0 updated" in resp.data
    mohan = Party.query.filter_by(name="Mohan").one()
    assert (mohan.phone, mohan.balance) == ("9876500001", Decimal("1200.00"))
    assert Party.query.filter_by(name="Sohan").one().balance == Decimal("-300.00")
    wa = client.get(f"/parties/{mohan.id}/remind").headers["Location"]
    assert wa.startswith("https://api.whatsapp.com/send?text=") and "phone=919876500001" in wa and "1%2C200.00" in wa
    assert b"WhatsApp reminder" in client.get("/reports/dues").data


# --- Direct printing ------------------------------------------------------------------------------------


def test_network_printer_receives_receipt_and_drawer_kick(client):
    server = socket.socket()
    server.bind(("127.0.0.1", 0))
    server.listen(1)
    received = []

    def accept():
        conn, _ = server.accept()
        with conn:
            chunks = iter(lambda: conn.recv(65536), b"")
            received.append(b"".join(chunks))

    thread = threading.Thread(target=accept, daemon=True)
    thread.start()
    client.post("/settings/devices", data={"printer_mode": "network", "printer_width": "48", "auto_print": "1",
                                           "open_drawer": "1",
                                           "printer_address": f"127.0.0.1:{server.getsockname()[1]}"})
    sale_id = new_sale(client)
    client.post(f"/sales/{sale_id}/scan", data={"code": SOAP})
    resp = client.post(f"/sales/{sale_id}/update", data={"action": "print"}, follow_redirects=True)
    thread.join(5)
    server.close()
    assert b"saved and printed" in resp.data
    data = received[0]
    assert data.startswith(b"\x1bp\x00") and b"P MART" in data and b"Medimix Soap 125g" in data
    assert data.endswith(b"\x1dVB\x03") and b"Rs.42.00" in data


def test_printer_unreachable_falls_back_to_screen(client):
    client.post("/settings/devices", data={"printer_mode": "network", "printer_address": "127.0.0.1:1",
                                           "printer_width": "48", "auto_print": "1"})
    sale_id = new_sale(client)
    client.post(f"/sales/{sale_id}/scan", data={"code": SOAP})
    resp = client.post(f"/sales/{sale_id}/update", data={"action": "print"})
    assert "/receipt" in resp.headers["Location"]
    assert db.session.get(Sale, sale_id).status == "saved"


# --- Export all and Google Drive -----------------------------------------------------------------------


def test_export_all_data(client):
    sell(client, SOAP)
    resp = client.get("/settings/export-all")
    wb = load_workbook(io.BytesIO(resp.data))
    assert "Products" in wb.sheetnames and "Bills & credit notes" in wb.sheetnames
    assert "password_hash" not in json.dumps([list(ws.values)[0] for ws in wb.worksheets])
    assert len(list(wb["Products"].values)) == 4


def test_google_drive_connect_and_upload(client, tmp_path, monkeypatch):
    import billing.drive as drive

    calls = []

    def fake_call(req):
        calls.append((req.get_method(), req.full_url))
        if req.full_url == drive.TOKEN_URL:
            body = req.data.decode()
            return {"refresh_token": "r-1"} if "authorization_code" in body else {"access_token": "a-1"}
        if req.full_url == drive.FILES_URL:
            return {"id": "folder-1"}
        return {"id": "file-1", "name": "x"}

    monkeypatch.setattr(drive, "_call", fake_call)
    monkeypatch.setattr(drive, "write_backup", lambda app, folder, keep: _fake_backup(tmp_path))
    client.post("/settings/drive", data={"drive_client_id": "cid.apps.googleusercontent.com",
                                         "drive_client_secret": "secret"})
    resp = client.post("/settings/drive/connect")
    assert resp.headers["Location"].startswith(drive.AUTH_URL) and "drive.file" in resp.headers["Location"]
    state = resp.headers["Location"].split("state=")[1]
    assert b"expired" in client.get("/settings/drive/callback?code=c&state=wrong", follow_redirects=True).data
    resp = client.post("/settings/drive/connect")
    state = resp.headers["Location"].split("state=")[1]
    assert b"Google Drive connected" in client.get(f"/settings/drive/callback?code=c&state={state}",
                                                   follow_redirects=True).data
    assert Setting.get().drive_refresh_token == "r-1"
    resp = client.post("/settings/drive/upload", follow_redirects=True)
    assert b"Uploaded" in resp.data
    uploads = [url for method, url in calls if url.startswith("https://www.googleapis.com/upload")]
    assert len(uploads) == 2 and Setting.get().drive_folder_id == "folder-1"


def _fake_backup(tmp_path):
    path = tmp_path / "billing-backup-test.db"
    path.write_bytes(b"SQLite format 3\x00")
    return str(path)


def test_drive_error_is_reported(client):
    Setting.get().drive_refresh_token = "r"
    Setting.get().drive_client_id = "c"
    db.session.commit()
    import billing.drive as drive

    def boom(req):
        raise drive.DriveError("Google said 401: invalid_grant")

    drive_call, drive._call = drive._call, boom
    try:
        resp = client.post("/settings/drive/upload", follow_redirects=True)
    finally:
        drive._call = drive_call
    assert b"Upload FAILED: Google said 401" in resp.data


def test_settings_sections_and_admin_only(app, client):
    for section in ("store", "gst", "devices", "backup", "audit"):
        assert client.get(f"/settings/{section}").status_code == 200, section
    client.post("/users/new", data={"name": "Manager", "username": "mgr", "password": "manage123",
                                    "perm_settings": "1", "active": "1"})
    other = app.test_client()
    other.post("/login", data={"username": "mgr", "password": "manage123"})
    assert other.get("/settings/devices").status_code == 200
    assert other.get("/settings/backup").status_code == 403
    assert other.get("/settings/export-all").status_code == 403
    assert User.query.filter_by(username="mgr").one().max_discount == Decimal("100")


def test_all_pages_render(client):
    sale = sell(client, SOAP)
    for url in ["/cash/", "/offers/", "/offers/new?kind=bxgy", "/offers/new?kind=bill", "/products/labels",
                "/purchases/reorder", "/customers/import", "/suppliers/import", f"/sales/{sale.id}/return",
                f"/sales/{sale.id}/invoice", "/purchases/orders/new"]:
        resp = client.get(url, follow_redirects=True)
        assert resp.status_code == 200, url


# --- Dashboard ----------------------------------------------------------------------------


def test_dashboard_periods_and_charts(client):
    sell(client, f"2*{SOAP}")
    for period in ["day", "week", "month", "year", "bogus"]:
        page = client.get(f"/reports/?period={period}").data.decode()
        assert 'data-chart="area"' in page and "Payment modes" in page, period
    page = client.get("/reports/?period=day").data.decode()
    assert "₹84" in page and "Received" in page
    old = client.get("/reports/?period=month&on=2020-02-10").data.decode()
    assert "February 2020" in old and "No sales in this period" in old and "Today" in old


def test_period_ranges():
    from billing.insights import Period

    week = Period.around("week", date(2026, 10, 25))  # a Sunday
    assert (week.start, week.end) == (date(2026, 10, 19), date(2026, 10, 26))
    fy = Period.around("year", date(2026, 2, 1))
    assert fy.label == "FY 2025-26" and len(fy.buckets()) == 12 and fy.buckets()[0][1] == "Apr"
    month = Period.around("month", date(2024, 2, 29))
    assert len(month.buckets()) == 29 and month.previous.label == "January 2024" and month.next.label == "March 2024"
    assert Period.around("day", date(2026, 1, 1)).buckets()[13][2] == "1pm – 2pm"
